"""Write reconciliation results to Excel.

`write_report` covers one folder, in the order someone reads it: the verdict first, then what
disagrees, then what could not be read, then the full matched set as evidence.

`write_overview` covers a date range -- one flat row per source per day, so a fortnight can be
scanned at a glance.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import DEFAULT_LABELS, KNOWN_FIELDS
from .matching import Result
from .records import DOCUMENT, Record

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="404040")
_GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")
_BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
_MAX_WIDTH = 60


def _write_header(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _autosize(sheet: Worksheet) -> None:
    for column in sheet.columns:
        width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
            max(width + 2, 10), _MAX_WIDTH
        )


def _fields_in_play(result: Result) -> list[str]:
    """Known fields that any record actually carries, match-key fields first."""
    present = {
        name
        for record in (*result.records_for(DOCUMENT), *result.unmatched)
        for name in record.values
    }
    ordered = [name for name in result.match_key if name in present]
    ordered += [n for n in KNOWN_FIELDS if n in present and n not in ordered]
    return ordered


def _summary_sheet(
    sheet: Worksheet, result: Result, source_name: str, folder: Path, labels: dict[str, str]
) -> None:
    sheet.title = "Summary"
    sheet.append(["Reconciliation summary"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    paired_by = (
        f"identifier: {result.match_on}"
        if result.match_on
        else f"composite key: {' + '.join(result.match_key)}"
    )
    rows = [
        ("Source", source_name),
        ("Folder", str(folder)),
        ("Records paired by", paired_by),
        ("", ""),
        (f"Records in {labels['document']}", result.document_count),
        (f"Records in {labels['spreadsheet']}", result.spreadsheet_count),
        ("Matched and agreeing", len(result.matched)),
    ]
    if result.match_on:
        rows.append(("Matched but differing", len(result.differing)))
    rows += [
        ("One-sided (unmatched)", len(result.unmatched)),
        ("Unread document lines", len(result.unread_lines)),
    ]
    if result.ambiguous_ids:
        rows.append(("Ambiguous identifiers", ", ".join(result.ambiguous_ids)))
    for warning in result.warnings:
        rows.append(("WARNING", warning))
    for name in ("amount", "quantity"):
        if name in result.match_key:
            doc_total, sheet_total = result.totals(name)
            rows.append((f"Total {name} (document)", round(doc_total, 6)))
            rows.append((f"Total {name} (spreadsheet)", round(sheet_total, 6)))
            rows.append((f"Difference in {name}", round(doc_total - sheet_total, 6)))
    for label, value in rows:
        sheet.append([label, value])
    for row in sheet.iter_rows(min_row=3, max_col=1):
        row[0].font = Font(bold=True)

    sheet.append([])
    sheet.append(["Verdict", result.verdict])
    verdict_row = sheet.max_row
    sheet.cell(verdict_row, 1).font = Font(bold=True)
    verdict_cell = sheet.cell(verdict_row, 2)
    verdict_cell.font = Font(bold=True)
    verdict_cell.fill = _GOOD_FILL if result.reconciled else _BAD_FILL
    _autosize(sheet)


def _record_row(record: Record, fields: list[str]) -> list[str]:
    return [record.display(name) for name in fields]


def _unmatched_sheet(
    workbook: Workbook, result: Result, fields: list[str], labels: dict[str, str]
) -> None:
    sheet = workbook.create_sheet("Unmatched")
    _write_header(sheet, ["Present in", "Missing from", "Where", *fields])
    for record in result.unmatched:
        from_document = record.origin == DOCUMENT
        present = labels["document"] if from_document else labels["spreadsheet"]
        missing = labels["spreadsheet"] if from_document else labels["document"]
        sheet.append([present, missing, record.locator, *_record_row(record, fields)])
    for row in sheet.iter_rows(min_row=2, max_col=1):
        row[0].fill = _BAD_FILL
    _autosize(sheet)


def _differences_sheet(workbook: Workbook, result: Result, labels: dict[str, str]) -> None:
    """One row per disagreeing field, so the sheet is a worklist rather than a puzzle."""
    sheet = workbook.create_sheet("Differences")
    _write_header(
        sheet,
        [
            result.match_on or "Identifier",
            "Field",
            f"In {labels['document']}",
            f"In {labels['spreadsheet']}",
            f"Where in {labels['document']}",
            f"Where in {labels['spreadsheet']}",
        ],
    )
    for pair in result.differing:
        for difference in pair.differences:
            sheet.append(
                [
                    pair.identifier,
                    difference.field,
                    difference.document_value,
                    difference.spreadsheet_value,
                    pair.document.locator,
                    pair.spreadsheet.locator,
                ]
            )
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].fill = _BAD_FILL
    _autosize(sheet)


def _unread_sheet(workbook: Workbook, result: Result) -> None:
    sheet = workbook.create_sheet("Unread Lines")
    _write_header(sheet, ["Document line the reader could not interpret"])
    for line in result.unread_lines:
        sheet.append([line])
    _autosize(sheet)


def _matched_sheet(
    workbook: Workbook, result: Result, fields: list[str], labels: dict[str, str]
) -> None:
    sheet = workbook.create_sheet("Matched")
    _write_header(
        sheet,
        [
            f"Where in {labels['document']}",
            f"Where in {labels['spreadsheet']}",
            *[f"{name}" for name in fields],
        ],
    )
    for pair in result.matched:
        sheet.append(
            [pair.document.locator, pair.spreadsheet.locator, *_record_row(pair.document, fields)]
        )
    _autosize(sheet)


def write_report(
    result: Result,
    out_path: str | Path,
    source_name: str,
    folder: Path,
    labels: dict[str, str] | None = None,
) -> Path:
    """Write the workbook and return where it landed.

    `labels` names the two sides in the output. Neither side has to be yours, so when both
    files arrive from outside these are what tell the reader which is which.
    """
    labels = {**DEFAULT_LABELS, **(labels or {})}
    out_path = Path(out_path)
    fields = _fields_in_play(result)

    workbook = Workbook()
    _summary_sheet(workbook.active, result, source_name, folder, labels)
    # Differences come first: a record that exists on both sides but disagrees is the most
    # actionable finding, and only id mode can produce one.
    if result.match_on:
        _differences_sheet(workbook, result, labels)
    _unmatched_sheet(workbook, result, fields, labels)
    _unread_sheet(workbook, result)
    _matched_sheet(workbook, result, fields, labels)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# Range overview: one row per source per day
# ---------------------------------------------------------------------------

_STATUS_FILL = {
    "RECONCILED": _GOOD_FILL,
    "DIFFERENCES": _BAD_FILL,
    "FAILED": _BAD_FILL,
    "INCOMPLETE": PatternFill("solid", fgColor="FFEB9C"),
}


def write_overview(outcomes, out_path: str | Path) -> Path:
    """Write the range overview workbook and return where it landed.

    Deliberately one flat sheet: the value is being able to scan a fortnight and see which cells
    are red, so anything that needs a second click defeats the purpose.
    """
    from .overview import tally

    out_path = Path(out_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"

    counts = tally(list(outcomes))
    sheet.append(["Reconciliation overview"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    summary_row = sheet.max_row + 1
    sheet.append(
        [
            f"{len(list(outcomes))} run(s)",
            f"{counts['RECONCILED']} reconciled",
            f"{counts['DIFFERENCES']} with differences",
            f"{counts['INCOMPLETE']} incomplete",
            f"{counts['FAILED']} failed",
        ]
    )
    for cell in sheet[summary_row]:
        cell.font = Font(bold=True)
    sheet.append([])

    header = [
        "Date",
        "Source",
        "Status",
        "Why",
        "Side A",
        "Side B",
        "Matched",
        "Differing",
        "One-sided",
        "Unread",
        "Folder",
    ]
    header_row = sheet.max_row + 1
    sheet.append(header)
    for cell in sheet[header_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    sheet.freeze_panes = sheet.cell(header_row + 1, 1)

    for outcome in outcomes:
        result = outcome.result
        sheet.append(
            [
                outcome.when.isoformat(),
                outcome.source,
                outcome.status,
                outcome.note,
                result.document_count if result else "",
                result.spreadsheet_count if result else "",
                len(result.matched) if result else "",
                len(result.differing) if result else "",
                len(result.unmatched) if result else "",
                len(result.unread_lines) if result else "",
                str(outcome.folder),
            ]
        )
        sheet.cell(sheet.max_row, 3).fill = _STATUS_FILL.get(outcome.status, _BAD_FILL)

    _autosize(sheet)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return out_path
