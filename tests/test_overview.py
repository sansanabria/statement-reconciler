"""Reconciling a whole date range in one pass."""

from __future__ import annotations

from datetime import date

import pytest
import yaml
from openpyxl import load_workbook

from statement_reconciler.cli import main
from statement_reconciler.overview import (
    DIFFERENCES,
    FAILED,
    INCOMPLETE,
    OVERVIEW_NAME,
    RECONCILED,
    run_range,
    tally,
)
from statement_reconciler.report import write_overview

from .conftest import ROWS

pytestmark = pytest.mark.integration

MARCH_2 = date(2026, 3, 2)
MARCH_3 = date(2026, 3, 3)


@pytest.fixture
def week(make_run_folder, config):
    """One clean day, and a second day whose spreadsheet is missing a record."""
    make_run_folder(ROWS, ROWS, name="EXAMPLE_SOURCE 2026.03.02")
    make_run_folder(ROWS, ROWS[:2], name="EXAMPLE_SOURCE 2026.03.03")
    return config


class TestRunRange:
    def test_each_day_gets_its_own_outcome(self, week):
        outcomes = run_range(week, [MARCH_2, MARCH_3])
        assert [o.status for o in outcomes] == [RECONCILED, DIFFERENCES]

    def test_the_reason_says_what_is_wrong(self, week):
        outcomes = run_range(week, [MARCH_2, MARCH_3])
        assert outcomes[0].note == ""
        assert outcomes[1].note == "1 one-sided"

    def test_a_missing_folder_is_incomplete_not_an_error(self, week):
        outcome = run_range(week, [date(2026, 3, 9)])[0]
        assert outcome.status == INCOMPLETE
        assert outcome.note == "folder does not exist"

    def test_a_half_filled_folder_names_the_gap(self, week, make_run_folder):
        folder = make_run_folder(ROWS, ROWS, name="EXAMPLE_SOURCE 2026.03.04")
        (folder / "records.xlsx").unlink()
        outcome = run_range(week, [date(2026, 3, 4)])[0]
        assert outcome.status == INCOMPLETE and outcome.note == "missing spreadsheet"

    def test_one_bad_folder_does_not_stop_the_range(self, week, make_run_folder):
        """A single unreadable day must not hide the state of every other day."""
        folder = make_run_folder(ROWS, ROWS, name="EXAMPLE_SOURCE 2026.03.04")
        (folder / "statement.pdf").write_bytes(b"not actually a pdf")
        outcomes = run_range(week, [MARCH_2, date(2026, 3, 4)])
        assert outcomes[0].status == RECONCILED
        assert outcomes[1].status == FAILED and outcomes[1].note

    def test_per_folder_reports_are_written(self, week):
        outcomes = run_range(week, [MARCH_2])
        assert outcomes[0].report_path is not None

    def test_no_report_mode_writes_nothing(self, week):
        outcomes = run_range(week, [MARCH_2], write_reports=False)
        assert outcomes[0].report_path is None

    def test_only_filters_by_source(self, week):
        assert run_range(week, [MARCH_2], only="ABSENT") == []


def test_tally_shows_zeroes_explicitly(week):
    counts = tally(run_range(week, [MARCH_2, MARCH_3]))
    assert counts == {RECONCILED: 1, DIFFERENCES: 1, INCOMPLETE: 0, FAILED: 0}


class TestOverviewWorkbook:
    def test_one_row_per_run(self, week, tmp_path):
        outcomes = run_range(week, [MARCH_2, MARCH_3])
        path = write_overview(outcomes, tmp_path / OVERVIEW_NAME)
        sheet = load_workbook(path)["Overview"]
        rows = [r for r in sheet.iter_rows(values_only=True) if r[0] == "2026-03-02"]
        assert len(rows) == 1

    def test_status_and_reason_are_on_the_row(self, week, tmp_path):
        outcomes = run_range(week, [MARCH_3])
        path = write_overview(outcomes, tmp_path / OVERVIEW_NAME)
        sheet = load_workbook(path)["Overview"]
        row = next(r for r in sheet.iter_rows(values_only=True) if r[0] == "2026-03-03")
        assert row[2] == DIFFERENCES and row[3] == "1 one-sided"

    def test_incomplete_rows_carry_no_counts(self, week, tmp_path):
        outcomes = run_range(week, [date(2026, 3, 9)])
        path = write_overview(outcomes, tmp_path / OVERVIEW_NAME)
        sheet = load_workbook(path)["Overview"]
        row = next(r for r in sheet.iter_rows(values_only=True) if r[0] == "2026-03-09")
        # openpyxl stores an empty string as an empty cell, so accept either.
        assert row[2] == INCOMPLETE and not row[4]


@pytest.fixture
def range_config_file(tmp_path, config_dict, week):
    config_dict["output_root"] = str(tmp_path)
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(config_dict), encoding="utf-8")
    return path


class TestRangeCli:
    def test_range_mode_reports_and_exits_one(self, range_config_file, capsys, tmp_path):
        code = main(
            [
                "check",
                "--config",
                str(range_config_file),
                "--from",
                "2026-03-02",
                "--to",
                "2026-03-03",
            ]
        )
        out = capsys.readouterr().out
        assert code == 1
        assert RECONCILED in out and DIFFERENCES in out
        assert (tmp_path / OVERVIEW_NAME).is_file()

    def test_a_clean_range_exits_zero(self, range_config_file, capsys):
        code = main(
            [
                "check",
                "--config",
                str(range_config_file),
                "--from",
                "2026-03-02",
                "--to",
                "2026-03-02",
            ]
        )
        assert code == 0

    def test_folder_and_range_together_is_rejected(self, range_config_file, tmp_path, capsys):
        code = main(
            [
                "check",
                "--config",
                str(range_config_file),
                str(tmp_path),
                "--from",
                "2026-03-02",
                "--to",
                "2026-03-03",
            ]
        )
        assert code == 2 and "not both" in capsys.readouterr().err

    def test_half_a_range_is_rejected(self, range_config_file, capsys):
        code = main(["check", "--config", str(range_config_file), "--from", "2026-03-02"])
        assert code == 2 and "both --from and --to" in capsys.readouterr().err

    def test_neither_folder_nor_range_is_rejected(self, range_config_file, capsys):
        code = main(["check", "--config", str(range_config_file)])
        assert code == 2 and "or --from/--to" in capsys.readouterr().err

    def test_single_folder_mode_still_works(self, range_config_file, run_folder, capsys):
        code = main(["check", "--config", str(range_config_file), str(run_folder)])
        assert code == 0 and "FULLY RECONCILED" in capsys.readouterr().out
