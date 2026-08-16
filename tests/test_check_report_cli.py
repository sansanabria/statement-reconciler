from __future__ import annotations

from datetime import date

import pytest
import yaml
from openpyxl import load_workbook

from statement_reconciler.check import REPORT_NAME, check_folder, find_files, infer_source
from statement_reconciler.cli import main
from statement_reconciler.config import ConfigError
from statement_reconciler.inventory import scan

from .conftest import ROWS

pytestmark = pytest.mark.integration


class TestFindFiles:
    def test_finds_one_of_each(self, run_folder, config):
        files = find_files(run_folder, config)
        assert files.document.suffix == ".pdf" and files.spreadsheet.suffix == ".xlsx"

    def test_two_pdfs_is_an_error_not_a_guess(self, run_folder, config):
        (run_folder / "extra.pdf").write_bytes(b"%PDF-1.4")
        with pytest.raises(ConfigError, match="2 PDF files found"):
            find_files(run_folder, config)

    def test_missing_side_is_named(self, run_folder, config):
        (run_folder / "records.xlsx").unlink()
        with pytest.raises(ConfigError, match="no spreadsheet found"):
            find_files(run_folder, config)

    def test_not_a_folder(self, tmp_path, config):
        with pytest.raises(ConfigError, match="Not a folder"):
            find_files(tmp_path / "absent", config)

    def test_an_existing_report_is_not_mistaken_for_input(self, run_folder, config):
        (run_folder / REPORT_NAME).write_bytes(b"PK")
        assert find_files(run_folder, config).spreadsheet.name == "records.xlsx"


class TestInferSource:
    def test_from_the_folder_path(self, run_folder, config):
        assert infer_source(run_folder, config).name == "EXAMPLE_SOURCE"

    def test_falls_back_to_the_only_source(self, tmp_path, config):
        folder = tmp_path / "somewhere else"
        folder.mkdir()
        assert infer_source(folder, config).name == "EXAMPLE_SOURCE"


class TestCheckFolder:
    def test_clean_folder_reconciles(self, run_folder, config):
        result, out_path = check_folder(run_folder, config)
        assert result.reconciled
        assert out_path == run_folder / REPORT_NAME and out_path.exists()

    def test_changed_quantity_becomes_two_one_sided_records(self, make_run_folder, config):
        altered = [(*ROWS[0][:2], "11", ROWS[0][3]), *ROWS[1:]]
        folder = make_run_folder(ROWS, altered)
        result, _ = check_folder(folder, config)
        assert not result.reconciled
        assert len(result.matched) == 2 and len(result.unmatched) == 2

    def test_row_missing_from_the_spreadsheet(self, make_run_folder, config):
        folder = make_run_folder(ROWS, ROWS[:2])
        result, _ = check_folder(folder, config)
        assert len(result.unmatched) == 1
        assert result.unmatched[0].values["description"] == "GAMMA DEVICE"

    def test_no_report_written_when_asked(self, run_folder, config):
        _, out_path = check_folder(run_folder, config, write=False)
        assert out_path is None and not (run_folder / REPORT_NAME).exists()


class TestReport:
    def test_workbook_has_the_four_sheets(self, run_folder, config):
        _, out_path = check_folder(run_folder, config)
        book = load_workbook(out_path)
        assert book.sheetnames == ["Summary", "Unmatched", "Unread Lines", "Matched"]

    def test_summary_states_the_verdict(self, run_folder, config):
        _, out_path = check_folder(run_folder, config)
        sheet = load_workbook(out_path)["Summary"]
        cells = [(row[0].value, row[1].value) for row in sheet.iter_rows(max_col=2)]
        assert ("Verdict", "FULLY RECONCILED") in cells
        assert ("Matched", 3) in cells

    def test_unmatched_sheet_says_which_side_is_missing_it(self, make_run_folder, config):
        folder = make_run_folder(ROWS, ROWS[:2])
        _, out_path = check_folder(folder, config)
        sheet = load_workbook(out_path)["Unmatched"]
        body = list(sheet.iter_rows(min_row=2, values_only=True))
        assert len(body) == 1
        assert body[0][0] == "Document" and body[0][1] == "Spreadsheet"

    def test_matched_sheet_lists_every_pair(self, run_folder, config):
        _, out_path = check_folder(run_folder, config)
        sheet = load_workbook(out_path)["Matched"]
        assert sheet.max_row == 4  # header plus three pairs


class TestInventory:
    def test_complete_and_incomplete_folders(self, run_folder, config):
        statuses = scan(config, [date(2026, 3, 2), date(2026, 3, 3)])
        by_date = {s.folder.name: s for s in statuses}
        assert by_date["EXAMPLE_SOURCE 2026.03.02"].complete
        missing = by_date["EXAMPLE_SOURCE 2026.03.03"]
        assert not missing.complete and missing.missing == "PDF + spreadsheet"

    def test_partial_folder_names_only_the_gap(self, run_folder, config):
        (run_folder / "records.xlsx").unlink()
        status = scan(config, [date(2026, 3, 2)])[0]
        assert status.missing == "spreadsheet"


@pytest.fixture
def config_file(tmp_path, config_dict, run_folder):
    config_dict["output_root"] = str(tmp_path)
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(config_dict), encoding="utf-8")
    return path


class TestCli:
    def test_check_returns_zero_when_reconciled(self, config_file, run_folder, capsys):
        code = main(["check", "--config", str(config_file), str(run_folder)])
        assert code == 0
        assert "FULLY RECONCILED" in capsys.readouterr().out

    def test_check_returns_one_on_differences(self, config_file, run_folder, capsys):
        (run_folder / "records.xlsx").unlink()
        (run_folder / "records.csv").write_text(
            "Date,Description,Quantity,Amount\n2026-03-02,ALPHA WIDGET,10,1234.50\n",
            encoding="utf-8",
        )
        assert main(["check", "--config", str(config_file), str(run_folder)]) == 1
        assert "DIFFERENCES FOUND" in capsys.readouterr().out

    def test_folders_creates_five_weekdays(self, config_file, tmp_path, config_dict):
        code = main(["folders", "--config", str(config_file), "--week", "2026-03-02"])
        assert code == 0

    def test_inventory_reports_gaps(self, config_file, capsys):
        code = main(
            [
                "inventory",
                "--config",
                str(config_file),
                "--from",
                "2026-03-02",
                "--to",
                "2026-03-04",
            ]
        )
        assert code == 1
        assert "MISSING" in capsys.readouterr().out

    def test_probe_needs_no_config(self, run_folder, capsys):
        assert main(["probe", str(run_folder / "statement.pdf"), "--mask"]) == 0
        assert "@" in capsys.readouterr().out

    def test_config_errors_are_reported_not_raised(self, tmp_path, capsys):
        code = main(["check", "--config", str(tmp_path / "absent.yaml"), str(tmp_path)])
        assert code == 2
        assert "Error:" in capsys.readouterr().err

    def test_bad_date_argument(self, config_file):
        with pytest.raises(SystemExit):
            main(["inventory", "--config", str(config_file), "--from", "yesterday", "--to", "x"])
